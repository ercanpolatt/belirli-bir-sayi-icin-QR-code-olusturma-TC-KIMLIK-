"""
QR Kod & Barkod Üretimi, Excel Entegrasyonu ve Pano Servisi
Desteklenen formatlar:
- QR Kod (2D)
- 1D Çizgi Barkodlar: Code 128, Code 39, EAN-13, EAN-8, UPC-A
- Kombine Kart (QR + Barkod)
Serbest metin, alfanümerik kod, seri numarası veya TC Kimlik Numarası ile çalışır.
"""
import os
import re
import csv
import io
import tempfile
from io import BytesIO
from PIL import Image, ImageDraw, ImageFont
import qrcode
import barcode
from barcode.writer import ImageWriter
import win32clipboard
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as OpenPyxlImage


def get_system_font(size: int, bold: bool = False) -> ImageFont.ImageFont:
    """
    Sistemde yüklü olan en uygun Türkçe destekli yazı tipini (Segoe UI, Arial, Calibri) döndürür.
    Bulamazsa varsayılan PIL fontunu yükler.
    """
    candidates = [
        "segoeuib.ttf" if bold else "segoeui.ttf",
        "arialbd.ttf" if bold else "arial.ttf",
        "calibrib.ttf" if bold else "calibri.ttf",
        "tahomabd.ttf" if bold else "tahoma.ttf",
    ]
    for font_name in candidates:
        try:
            return ImageFont.truetype(font_name, size)
        except IOError:
            continue
    return ImageFont.load_default()


def copy_image_to_clipboard(pil_image: Image.Image) -> bool:
    """
    Verilen PIL Image görselini doğrudan Windows panosuna (Clipboard) kopyalar.
    Böylece kullanıcı Excel, Word veya Paint'e geçip Ctrl+V ile görseli doğrudan yapıştırabilir.
    """
    if pil_image is None:
        return False
    try:
        output = io.BytesIO()
        pil_image.convert("RGB").save(output, "BMP")
        data = output.getvalue()[14:]  # 14 byte BMP dosya başlığını çıkarıp CF_DIB formatı oluştur
        output.close()
        
        win32clipboard.OpenClipboard()
        win32clipboard.EmptyClipboard()
        win32clipboard.SetClipboardData(win32clipboard.CF_DIB, data)
        win32clipboard.CloseClipboard()
        return True
    except Exception as e:
        print(f"Pano görsel kopyalama hatası: {e}")
        return False


def copy_files_to_clipboard_as_hdrop(images: list[Image.Image]) -> bool:
    """
    Verilen PIL Image listesini geçici dosyalara kaydeder ve Windows panosuna (CF_HDROP) kopyalar.
    Word vb. programlara Ctrl+V ile doğrudan çoklu resim (dosya) olarak yapıştırılmasını sağlar.
    """
    import struct
    if not images:
        return False
        
    try:
        temp_dir = tempfile.mkdtemp()
        file_paths = []
        for idx, img in enumerate(images):
            path = os.path.abspath(os.path.join(temp_dir, f"Sayfa_{idx+1}.png"))
            img.save(path, format="PNG")
            file_paths.append(path)
            
        paths_str = "\0".join(file_paths) + "\0\0"
        encoded_paths = paths_str.encode("utf-16le")
        dropfiles = struct.pack("IIIII", 20, 0, 0, 0, 1)
        data = dropfiles + encoded_paths
        
        win32clipboard.OpenClipboard()
        win32clipboard.EmptyClipboard()
        win32clipboard.SetClipboardData(win32clipboard.CF_HDROP, data)
        win32clipboard.CloseClipboard()
        return True
    except Exception as e:
        print(f"Pano CF_HDROP hatası: {e}")
        return False


def sanitize_for_barcode(text: str, barcode_type: str = "code128") -> tuple[str, str]:
    """
    Barkod üretimi için metni normalize eder ve seçilen barkod tipine uygunluğunu denetler.
    Dönüş: (temizlenmis_metin, hata_mesaji)
    """
    if not text:
        return "", "Kod değeri boş olamaz."
    text = str(text).strip()
    
    # Türkçe karakter dönüşüm haritası
    tr_map = str.maketrans("çğıöşüÇĞİÖŞÜ", "cgiosuCGIOSU")
    clean = text.translate(tr_map)
    
    type_key = barcode_type.lower().replace("-", "").replace(" ", "").replace("_", "")
    
    if "ean13" in type_key:
        digits = "".join(c for c in clean if c.isdigit())
        if len(digits) not in [12, 13]:
            return "", f"EAN-13 için 12 veya 13 haneli rakam gereklidir (Girilen: {len(digits)} hane)."
        return digits[:12], ""  # python-barcode 12 hane alıp checksum'ı kendisi hesaplar
        
    elif "ean8" in type_key:
        digits = "".join(c for c in clean if c.isdigit())
        if len(digits) not in [7, 8]:
            return "", f"EAN-8 için 7 veya 8 haneli rakam gereklidir (Girilen: {len(digits)} hane)."
        return digits[:7], ""
        
    elif "upca" in type_key or "upc" in type_key:
        digits = "".join(c for c in clean if c.isdigit())
        if len(digits) not in [11, 12]:
            return "", f"UPC-A için 11 veya 12 haneli rakam gereklidir (Girilen: {len(digits)} hane)."
        return digits[:11], ""
        
    elif "39" in type_key:
        clean = clean.upper()
        allowed = set("0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ-. $/+%")
        filtered = "".join(c for c in clean if c in allowed)
        if not filtered:
            return "", "Code 39 formatı için geçerli karakter bulunamadı (Harf, rakam, -. $/+%)."
        return filtered, ""
        
    else:  # code128 varsayılan
        # ASCII karakter kontrolü
        filtered = "".join(c for c in clean if ord(c) < 128)
        if not filtered:
            return "", "Code 128 için geçerli ASCII karakteri bulunamadı."
        return filtered, ""


def validate_tckn(tckn: str) -> tuple[bool, str]:
    """
    TC Kimlik Numarasının matematiksel geçerliliğini kontrol eder.
    """
    if not tckn:
        return False, "TC Kimlik numarası boş olamaz."
    
    tckn = str(tckn).strip()
    
    if len(tckn) != 11:
        return False, "TC Kimlik numarası 11 haneli olmalıdır."
    
    if not tckn.isdigit():
        return False, "TC Kimlik numarası sadece rakamlardan oluşmalıdır."
    
    if tckn[0] == '0':
        return False, "TC Kimlik numarası 0 ile başlayamaz."
    
    digits = [int(d) for d in tckn]
    
    odd_sum = sum(digits[0:9:2])
    even_sum = sum(digits[1:8:2])
    digit10 = ((odd_sum * 7) - even_sum) % 10
    
    if digits[9] != digit10:
        return False, "Geçersiz TC Kimlik Numarası (Algoritma kontrolü)."
    
    digit11 = sum(digits[:10]) % 10
    if digits[10] != digit11:
        return False, "Geçersiz TC Kimlik Numarası (Toplam kontrolü)."
    
    return True, "Geçerli TC Kimlik Numarası."


def validate_code_value(code_val: str, strict_tckn: bool = False, barcode_type: str = "code128") -> tuple[bool, str]:
    """
    Girilen kod değerini doğrular.
    - strict_tckn=True ise zorunlu 11 haneli TC kontrolü yapar.
    - strict_tckn=False ise serbest sayı, harf, seri no veya seçilen barkod tipini doğrular.
    """
    if not code_val or not str(code_val).strip():
        return False, "Barkod / Kod değeri boş olamaz."
        
    code_val = str(code_val).strip()
    
    if strict_tckn:
        return validate_tckn(code_val)
        
    # Barkod türüne özel kontrol
    clean_val, err = sanitize_for_barcode(code_val, barcode_type=barcode_type)
    if err:
        return False, err
        
    if len(code_val) == 11 and code_val.isdigit():
        is_tckn, _ = validate_tckn(code_val)
        if is_tckn:
            return True, "Geçerli TC Kimlik Numarası algılandı."
        else:
            return True, f"Özel Kod / Sayı Değeri ({len(code_val)} Karakter)"
            
    return True, f"Geçerli Kod Değeri ({len(clean_val)} Karakter)"


def generate_qr_image(data: str, box_size: int = 10, border: int = 2) -> Image.Image:
    """
    Girilen metin, sayı veya seri numarasını içeren QR kod görseli oluşturur.
    """
    data_str = str(data).strip()
    qr = qrcode.QRCode(
        version=1,
        error_correction=qrcode.constants.ERROR_CORRECT_M,
        box_size=box_size,
        border=border,
    )
    qr.add_data(data_str)
    qr.make(fit=True)
    
    img = qr.make_image(fill_color="black", back_color="white")
    return img.convert("RGB")


def generate_qr_with_text(data: str, name: str = "", surname: str = "", box_size: int = 12, border: int = 2, show_text: bool = True) -> Image.Image:
    """
    QR kod görseli oluşturur ve altına ismini ve istenirse TC (veya Kod) değerini yazar.
    """
    qr_img = generate_qr_image(data, box_size=box_size, border=border)
    
    font_name = get_system_font(26, bold=True)
    font_code = get_system_font(18, bold=False)
    
    full_name = f"{str(name).upper()} {str(surname).upper()}".strip()
    code_str = str(data).strip()
    
    extra_height = 80 if full_name else 40
    new_width = max(qr_img.width, 300)
    new_height = qr_img.height + extra_height
    
    combined = Image.new("RGB", (new_width, new_height), "white")
    
    qr_x = (new_width - qr_img.width) // 2
    combined.paste(qr_img, (qr_x, 0))
    
    draw = ImageDraw.Draw(combined)
    current_y = qr_img.height + 5
    
    if full_name:
        bbox = draw.textbbox((0, 0), full_name, font=font_name)
        text_w = bbox[2] - bbox[0]
        draw.text(((new_width - text_w) // 2, current_y), full_name, fill="#0F172A", font=font_name)
        current_y += 35
        
    if show_text:
        lbl_str = f"TC/KOD: {code_str}"
        bbox_lbl = draw.textbbox((0, 0), lbl_str, font=font_code)
        lbl_w = bbox_lbl[2] - bbox_lbl[0]
        draw.text(((new_width - lbl_w) // 2, current_y), lbl_str, fill="#475569", font=font_code)
    
    return combined


def generate_barcode_image(
    data: str,
    barcode_type: str = "code128",
    show_text: bool = True,
    module_width: float = 0.42,
    module_height: float = 16.0,
    font_size: int = 14
) -> Image.Image:
    """
    Herhangi bir sayı, harf veya seri numarası için 1D çizgi barkod üretir.
    - barcode_type: "code128", "code39", "ean13", "ean8", "upca"
    """
    raw_str = str(data).strip()
    clean_str, err = sanitize_for_barcode(raw_str, barcode_type=barcode_type)
    if not clean_str:
        clean_str = "0"
        
    type_key = barcode_type.lower().replace("-", "").replace(" ", "").replace("_", "")
    
    if "ean13" in type_key:
        bc_class_name = "ean13"
    elif "ean8" in type_key:
        bc_class_name = "ean8"
    elif "upca" in type_key or "upc" in type_key:
        bc_class_name = "upca"
    elif "39" in type_key:
        bc_class_name = "code39"
    else:
        bc_class_name = "code128"
        
    try:
        bc_cls = barcode.get_barcode_class(bc_class_name)
        writer = ImageWriter()
        
        writer_options = {
            "module_width": module_width,
            "module_height": module_height,
            "font_size": font_size,
            "text_distance": 4.0,
            "quiet_zone": 4.0,
            "write_text": show_text
        }
        
        if bc_class_name == "code39":
            bc = bc_cls(clean_str, writer=writer, add_checksum=False)
        else:
            bc = bc_cls(clean_str, writer=writer)
            
        fp = BytesIO()
        bc.write(fp, options=writer_options)
        fp.seek(0)
        img = Image.open(fp).convert("RGB")
        return img
    except Exception as e:
        # Hata durumunda fallback olarak Code 128 dene
        try:
            bc_cls = barcode.get_barcode_class("code128")
            writer = ImageWriter()
            bc = bc_cls(str(data).strip() or "0", writer=writer)
            fp = BytesIO()
            bc.write(fp, options={"module_width": module_width, "module_height": module_height, "write_text": show_text})
            fp.seek(0)
            return Image.open(fp).convert("RGB")
        except Exception:
            # En son çare boş beyaz resim
            return Image.new("RGB", (300, 100), "white")


def generate_combined_code_image(
    data: str,
    barcode_type: str = "code128",
    show_barcode_text: bool = True
) -> Image.Image:
    """
    Hem QR Kod hem de Çizgi Barkodu tek bir görselde estetik olarak birleştirir.
    """
    data_str = str(data).strip()
    qr_img = generate_qr_image(data_str, box_size=8, border=2)
    bc_img = generate_barcode_image(data_str, barcode_type=barcode_type, show_text=show_barcode_text, module_width=0.38, module_height=14.0)
    
    total_w = max(qr_img.width + 40, bc_img.width + 40, 600)
    total_h = qr_img.height + bc_img.height + 60
    
    combined = Image.new("RGB", (total_w, total_h), "white")
    
    qr_x = (total_w - qr_img.width) // 2
    qr_y = 20
    combined.paste(qr_img, (qr_x, qr_y))
    
    bc_x = (total_w - bc_img.width) // 2
    bc_y = qr_y + qr_img.height + 20
    combined.paste(bc_img, (bc_x, bc_y))
    
    return combined


def create_printable_badge(
    name: str,
    surname: str,
    code_value: str,
    code_mode: str = "barcode",
    barcode_type: str = "code128",
    show_barcode_text: bool = True,
    company_title: str = ""
) -> Image.Image:
    """
    Yazıcı çıktısı ve etiket için yüksek kaliteli, estetik kart görseli oluşturur (800x850 px).
    - code_mode: "qr", "barcode", "both"
    """
    width, height = 800, 850
    card = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(card)
    
    font_name = get_system_font(40, bold=True)
    font_code = get_system_font(26, bold=False)
    font_header = get_system_font(25, bold=True)
    
    code_str = str(code_value).strip()
    full_name = f"{str(name).upper()} {str(surname).upper()}".strip()
    has_name = bool(full_name and full_name != "-")
    
    start_y = 35
    
    # Opsiyonel Firma / Kurum Başlığı
    if company_title.strip():
        comp_text = company_title.strip().upper()
        bbox_comp = draw.textbbox((0, 0), comp_text, font=font_header)
        comp_w = bbox_comp[2] - bbox_comp[0]
        draw.text(((width - comp_w) // 2, start_y), comp_text, fill="#1E293B", font=font_header)
        draw.line([(50, start_y + 40), (width - 50, start_y + 40)], fill="#E2E8F0", width=2)
        start_y += 55

    label_prefix = "TCKN" if (len(code_str) == 11 and code_str.isdigit()) else "KOD"

    # 1. YALNIZCA QR KOD MODU
    if code_mode == "qr":
        qr_img = generate_qr_image(code_str, box_size=15, border=2)
        qr_w, qr_h = qr_img.size
        qr_x = (width - qr_w) // 2
        qr_y = start_y + (10 if has_name else 40)
        card.paste(qr_img, (qr_x, qr_y))
        
        current_y = qr_y + qr_h + 25
        if has_name:
            bbox = draw.textbbox((0, 0), full_name, font=font_name)
            text_w = bbox[2] - bbox[0]
            draw.text(((width - text_w) // 2, current_y), full_name, fill="#0F172A", font=font_name)
            current_y += 55
            
        if show_barcode_text:
            lbl_str = f"{label_prefix}: {code_str}"
            bbox_lbl = draw.textbbox((0, 0), lbl_str, font=font_code)
            lbl_w = bbox_lbl[2] - bbox_lbl[0]
            draw.text(((width - lbl_w) // 2, current_y), lbl_str, fill="#475569", font=font_code)

    # 2. YALNIZCA BARKOD MODU
    elif code_mode == "barcode":
        bc_img = generate_barcode_image(
            code_str,
            barcode_type=barcode_type,
            show_text=show_barcode_text,
            module_width=0.48,
            module_height=26.0,
            font_size=18
        )
        
        if bc_img.width > (width - 60):
            ratio = (width - 60) / bc_img.width
            new_h = int(bc_img.height * ratio)
            bc_img = bc_img.resize((width - 60, new_h), Image.Resampling.LANCZOS)
            
        bc_w, bc_h = bc_img.size
        bc_x = (width - bc_w) // 2
        bc_y = start_y + (60 if has_name else 100)
        card.paste(bc_img, (bc_x, bc_y))
        
        current_y = bc_y + bc_h + 40
        if has_name:
            bbox = draw.textbbox((0, 0), full_name, font=font_name)
            text_w = bbox[2] - bbox[0]
            draw.text(((width - text_w) // 2, current_y), full_name, fill="#0F172A", font=font_name)
            current_y += 55
            
        if not show_barcode_text or not has_name:
            lbl_str = f"{label_prefix}: {code_str}"
            bbox_lbl = draw.textbbox((0, 0), lbl_str, font=font_code)
            lbl_w = bbox_lbl[2] - bbox_lbl[0]
            draw.text(((width - lbl_w) // 2, current_y), lbl_str, fill="#475569", font=font_code)

    # 3. KOMBİNE MOD (QR KOD + BARKOD BİRLİKTE)
    else:
        qr_img = generate_qr_image(code_str, box_size=9, border=2)
        qr_w, qr_h = qr_img.size
        qr_x = (width - qr_w) // 2
        qr_y = start_y + 5
        card.paste(qr_img, (qr_x, qr_y))
        
        bc_img = generate_barcode_image(
            code_str,
            barcode_type=barcode_type,
            show_text=show_barcode_text,
            module_width=0.40,
            module_height=14.0,
            font_size=15
        )
        if bc_img.width > (width - 80):
            ratio = (width - 80) / bc_img.width
            new_h = int(bc_img.height * ratio)
            bc_img = bc_img.resize((width - 80, new_h), Image.Resampling.LANCZOS)
            
        bc_w, bc_h = bc_img.size
        bc_x = (width - bc_w) // 2
        bc_y = qr_y + qr_h + 18
        card.paste(bc_img, (bc_x, bc_y))
        
        current_y = bc_y + bc_h + 20
        if has_name:
            bbox = draw.textbbox((0, 0), full_name, font=font_name)
            text_w = bbox[2] - bbox[0]
            draw.text(((width - text_w) // 2, current_y), full_name, fill="#0F172A", font=font_name)
        else:
            lbl_str = f"{label_prefix}: {code_str}"
            bbox_lbl = draw.textbbox((0, 0), lbl_str, font=font_code)
            lbl_w = bbox_lbl[2] - bbox_lbl[0]
            draw.text(((width - lbl_w) // 2, current_y), lbl_str, fill="#475569", font=font_code)

    return card


def create_grid_printable_pages(
    workers: list[tuple[str, str, str]],
    items_per_row: int = 3,
    rows_per_page: int = 4,
    code_mode: str = "barcode",
    barcode_type: str = "code128",
    show_barcode_text: bool = True,
    company_title: str = ""
) -> list[Image.Image]:
    """
    A4 sayfasına yan yana 3'lü dizilimde yazdırılabilir sayfalar oluşturur (2480 x 3508 px - 300 DPI).
    """
    if not workers:
        return []
        
    page_width, page_height = 2480, 3508
    margin_x = 115
    margin_y = 100
    gap_x = 30
    gap_y = 40
    
    cell_w = 740
    cell_h = 770
    
    capacity_per_page = items_per_row * rows_per_page
    pages = []
    
    font_name = get_system_font(56, bold=True)
    font_tc = get_system_font(36, bold=False)

    current_page = None
    draw = None
    
    for idx, (name, surname, code_val) in enumerate(workers):
        pos_in_page = idx % capacity_per_page
        
        if pos_in_page == 0:
            current_page = Image.new("RGB", (page_width, page_height), "white")
            draw = ImageDraw.Draw(current_page)
            pages.append(current_page)
            
        row = pos_in_page // items_per_row
        col = pos_in_page % items_per_row
        
        x = margin_x + col * (cell_w + gap_x)
        y = margin_y + row * (cell_h + gap_y)
        
        draw.rectangle([(x, y), (x + cell_w, y + cell_h)], outline="#CBD5E1", width=2)
        
        code_str = str(code_val).strip()
        full_name = f"{str(name).upper()} {str(surname).upper()}".strip()
        has_name = bool(full_name and full_name != "-")
        label_prefix = "TC" if (len(code_str) == 11 and code_str.isdigit()) else "KOD"

        if code_mode == "qr":
            qr_img = generate_qr_image(code_str, box_size=20, border=2)
            qr_w, qr_h = qr_img.size
            qr_x = x + (cell_w - qr_w) // 2
            qr_y = y + (50 if has_name else 120)
            current_page.paste(qr_img, (qr_x, qr_y))
            
            cur_y = qr_y + qr_h + 30
            if has_name:
                bbox = draw.textbbox((0, 0), full_name, font=font_name)
                text_w = bbox[2] - bbox[0]
                draw.text((x + (cell_w - text_w) // 2, cur_y), full_name, fill="#000000", font=font_name)
                cur_y += 75
                
            if show_barcode_text:
                tc_str = f"{label_prefix}: {code_str}"
                bbox_tc = draw.textbbox((0, 0), tc_str, font=font_tc)
                tc_w = bbox_tc[2] - bbox_tc[0]
                draw.text((x + (cell_w - tc_w) // 2, cur_y), tc_str, fill="#64748B", font=font_tc)

        elif code_mode == "barcode":
            bc_img = generate_barcode_image(
                code_str,
                barcode_type=barcode_type,
                show_text=show_barcode_text,
                module_width=0.75,
                module_height=45.0,
                font_size=26
            )
            if bc_img.width > (cell_w - 60):
                ratio = (cell_w - 60) / bc_img.width
                bc_img = bc_img.resize((cell_w - 60, int(bc_img.height * ratio)), Image.Resampling.LANCZOS)
                
            bc_w, bc_h = bc_img.size
            bc_x = x + (cell_w - bc_w) // 2
            bc_y = y + (160 if has_name else 220)
            current_page.paste(bc_img, (bc_x, bc_y))
            
            cur_y = bc_y + bc_h + 40
            if has_name:
                bbox = draw.textbbox((0, 0), full_name, font=font_name)
                text_w = bbox[2] - bbox[0]
                draw.text((x + (cell_w - text_w) // 2, cur_y), full_name, fill="#000000", font=font_name)
            else:
                if show_barcode_text:
                    tc_str = f"{label_prefix}: {code_str}"
                    bbox_tc = draw.textbbox((0, 0), tc_str, font=font_tc)
                    tc_w = bbox_tc[2] - bbox_tc[0]
                    draw.text((x + (cell_w - tc_w) // 2, cur_y), tc_str, fill="#64748B", font=font_tc)

        else:
            qr_img = generate_qr_image(code_str, box_size=15, border=2)
            qr_w, qr_h = qr_img.size
            qr_x = x + (cell_w - qr_w) // 2
            qr_y = y + 40
            current_page.paste(qr_img, (qr_x, qr_y))
            
            bc_img = generate_barcode_image(
                code_str,
                barcode_type=barcode_type,
                show_text=show_barcode_text,
                module_width=0.6,
                module_height=25.0,
                font_size=20
            )
            if bc_img.width > (cell_w - 60):
                ratio = (cell_w - 60) / bc_img.width
                bc_img = bc_img.resize((cell_w - 60, int(bc_img.height * ratio)), Image.Resampling.LANCZOS)
                
            bc_w, bc_h = bc_img.size
            bc_x = x + (cell_w - bc_w) // 2
            bc_y = qr_y + qr_h + 20
            current_page.paste(bc_img, (bc_x, bc_y))
            
            cur_y = bc_y + bc_h + 30
            if has_name:
                bbox = draw.textbbox((0, 0), full_name, font=font_name)
                text_w = bbox[2] - bbox[0]
                draw.text((x + (cell_w - text_w) // 2, cur_y), full_name, fill="#000000", font=font_name)
            else:
                if show_barcode_text:
                    tc_str = f"{label_prefix}: {code_str}"
                    bbox_tc = draw.textbbox((0, 0), tc_str, font=font_tc)
                    tc_w = bbox_tc[2] - bbox_tc[0]
                    draw.text((x + (cell_w - tc_w) // 2, cur_y), tc_str, fill="#64748B", font=font_tc)

    return pages


def smart_extract_tckn(cell_value) -> str:
    """
    Hücre değerinden 11 haneli TC Kimlik numarasını veya varsa alfanümerik kod değerini çıkarır.
    """
    if cell_value is None:
        return ""
        
    s = str(cell_value).strip()
    if not s:
        return ""
        
    if isinstance(cell_value, (float, int)):
        try:
            s = f"{cell_value:.0f}"
        except Exception:
            pass
    elif s.endswith(".0"):
        s = s[:-2]
        
    clean_digits = "".join([c for c in s if c.isdigit()])
    
    if len(clean_digits) == 11 and clean_digits[0] != '0':
        return clean_digits
            
    matches = re.findall(r'\d{11}', s)
    for m in matches:
        if m[0] != '0':
            return m
            
    return s


def smart_parse_excel_or_csv(filepath: str) -> list[dict]:
    """
    Excel / CSV dosyasındaki kayıtları (İsim, Soyisim ve Kod/TC No) akıllıca çıkarır.
    """
    all_sheet_rows = []
    
    if filepath.lower().endswith(".csv"):
        sheet_rows = []
        for encoding in ["utf-8-sig", "utf-8", "cp1254", "iso-8859-9", "latin-1"]:
            try:
                with open(filepath, mode="r", encoding=encoding) as f:
                    sample = f.read(4096)
                    f.seek(0)
                    dialect = csv.excel
                    try:
                        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|\t")
                    except Exception:
                        pass
                    reader = csv.reader(f, dialect)
                    for r in reader:
                        sheet_rows.append([str(cell).strip() if cell else "" for cell in r])
                if sheet_rows:
                    break
            except Exception:
                continue
        if sheet_rows:
            all_sheet_rows.append(sheet_rows)
    else:
        try:
            wb = openpyxl.load_workbook(filepath, data_only=True)
            for sheet in wb.worksheets:
                sheet_rows = []
                for row in sheet.iter_rows(values_only=True):
                    row_cells = []
                    for cell in row:
                        if cell is not None:
                            if isinstance(cell, (float, int)):
                                if isinstance(cell, float) and cell.is_integer():
                                    row_cells.append(str(int(cell)))
                                else:
                                    row_cells.append(f"{cell:.0f}" if isinstance(cell, float) else str(cell))
                            else:
                                row_cells.append(str(cell).strip())
                        else:
                            row_cells.append("")
                    if any(row_cells):
                        sheet_rows.append(row_cells)
                if sheet_rows:
                    all_sheet_rows.append(sheet_rows)
        except Exception as e:
            print(f"Excel Okuma Hatası: {e}")

    return parse_raw_grid_rows(all_sheet_rows)


def parse_clipboard_table_text(text: str) -> list[dict]:
    """
    Excel'den kopyalanmış sekmeli (TSV) veya virgüllü/noktalı virgüllü metinleri ayrıştırır.
    """
    if not text or not text.strip():
        return []
        
    lines = [line for line in text.strip().splitlines() if line.strip()]
    grid = []
    for line in lines:
        if "\t" in line:
            parts = [p.strip() for p in line.split("\t")]
        elif ";" in line:
            parts = [p.strip() for p in line.split(";")]
        elif "," in line:
            parts = [p.strip() for p in line.split(",")]
        else:
            parts = [line.strip()]
        grid.append(parts)
        
    return parse_raw_grid_rows([grid])


def parse_raw_grid_rows(all_sheet_rows: list[list[list[str]]]) -> list[dict]:
    """
    Ham satır/hücre listesinden İsim, Soyisim ve Kod/TC kayıtlarını üretir.
    """
    workers = []
    seen_codes = set()

    for raw_rows in all_sheet_rows:
        for row in raw_rows:
            if not any(row):
                continue
                
            code_found = ""
            code_col_idx = -1
            
            # 1. Öncelik: 11 haneli TCKN ara
            for col_idx, cell_val in enumerate(row):
                cand = smart_extract_tckn(cell_val)
                if len(cand) == 11 and cand.isdigit() and cand not in seen_codes:
                    code_found = cand
                    code_col_idx = col_idx
                    break
                    
            # 2. Öncelik: Herhangi bir belirgin kod/sayı/seri no ara
            if not code_found:
                for col_idx, cell_val in enumerate(row):
                    val_str = str(cell_val).strip()
                    if val_str and val_str not in seen_codes and len(val_str) >= 1:
                        val_upper = val_str.upper()
                        if val_upper in ["TC", "TCKN", "TC NO", "T.C.", "KOD", "BARKOD", "BARCODE", "AD", "SOYAD", "AD SOYAD", "İSİM", "SOYİSİM", "SIRA NO", "NO", "SIRA", "ID", "TARİH"]:
                            continue
                        code_found = val_str
                        code_col_idx = col_idx
                        break
                            
            if not code_found:
                continue
                
            text_candidates = []
            for col_idx, cell_val in enumerate(row):
                if col_idx == code_col_idx:
                    continue
                    
                val_str = str(cell_val).strip()
                if not val_str:
                    continue
                    
                clean_text = re.sub(r'^\d+[\s\.\-\)]*', '', val_str).strip()
                if not clean_text:
                    continue
                    
                if clean_text.isdigit() or clean_text.replace('.', '', 1).isdigit():
                    continue
                    
                val_upper = clean_text.upper()
                if val_upper in ["TC", "TCKN", "TC NO", "T.C.", "KOD", "BARKOD", "BARCODE", "AD", "SOYAD", "AD SOYAD", "İSİM", "SOYİSİM", "SIRA NO", "NO", "SIRA", "ID", "TARİH"]:
                    continue
                    
                text_candidates.append(clean_text)

            name = ""
            surname = ""
            
            multi_word_cells = [c for c in text_candidates if len(c.split()) >= 2]
            if multi_word_cells:
                target = multi_word_cells[0]
                parts = target.split()
                surname = parts[-1]
                name = " ".join(parts[:-1])
            elif len(text_candidates) >= 2:
                name = text_candidates[0]
                surname = text_candidates[1]
            elif len(text_candidates) == 1:
                name = text_candidates[0]
                surname = ""
            else:
                name = ""
                surname = ""

            workers.append({
                "name": name.strip(),
                "surname": surname.strip(),
                "tckn": code_found
            })
            seen_codes.add(code_found)

    return workers


def export_records_to_excel(
    records: list[tuple],
    filepath: str,
    include_images: bool = False,
    code_mode: str = "barcode",
    barcode_type: str = "code128"
):
    """
    Kayıt listesini şık biçimlendirilmiş bir .xlsx Excel tablosu olarak kaydeder.
    include_images=True ise her satırın yanına gerçek barkod görselini gömer.
    records: list of (id, name, surname, code_val, created_at)
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Barkod Listesi"
    
    headers = ["Sıra / ID", "İsim / Başlık", "Soyisim / Detay", "Barkod / Kod Değeri", "Kayıt Tarihi"]
    if include_images:
        headers.append("Barkod / QR Görseli")
        
    ws.append(headers)
    
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    
    ws.row_dimensions[1].height = 28
    
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    
    temp_dir = tempfile.mkdtemp()
    temp_images = []
    
    for r_idx, row in enumerate(records, start=2):
        if include_images:
            ws.row_dimensions[r_idx].height = 65
        else:
            ws.row_dimensions[r_idx].height = 22
            
        for c_idx, val in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=str(val))
            cell.font = Font(name="Segoe UI", size=10)
            cell.border = thin_border
            if c_idx in [1, 4, 5]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
                
        if include_images:
            # Görsel oluşturup hücreye yerleştir
            code_val = str(row[3])
            img_cell = ws.cell(row=r_idx, column=6)
            img_cell.border = thin_border
            img_cell.alignment = Alignment(horizontal="center", vertical="center")
            
            try:
                if code_mode == "qr":
                    pil_img = generate_qr_image(code_val, box_size=4, border=1)
                elif code_mode == "both":
                    pil_img = generate_combined_code_image(code_val, barcode_type=barcode_type, show_barcode_text=True)
                else:
                    pil_img = generate_barcode_image(code_val, barcode_type=barcode_type, show_text=True, module_width=0.25, module_height=10.0, font_size=10)
                
                # Boyutlandır (Hücreye sığdır)
                max_w, max_h = 160, 55
                pil_img.thumbnail((max_w, max_h), Image.Resampling.LANCZOS)
                
                temp_img_path = os.path.join(temp_dir, f"cell_img_{r_idx}.png")
                pil_img.save(temp_img_path)
                temp_images.append(temp_img_path)
                
                xl_img = OpenPyxlImage(temp_img_path)
                xl_img.width = pil_img.width
                xl_img.height = pil_img.height
                
                cell_coord = f"F{r_idx}"
                ws.add_image(xl_img, cell_coord)
            except Exception as img_err:
                print(f"Excel görsel ekleme hatası: {img_err}")
                
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        if col_letter == "F" and include_images:
            ws.column_dimensions[col_letter].width = 28
        else:
            max_len = max(len(str(cell.value or '')) for cell in col)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 15)
        
    wb.save(filepath)
    
    # Geçici dosyaları temizle
    for p in temp_images:
        try:
            os.remove(p)
        except Exception:
            pass
    try:
        os.rmdir(temp_dir)
    except Exception:
        pass


def generate_sequential_codes(
    start_num: int,
    count: int,
    prefix: str = "",
    suffix: str = "",
    pad_length: int = 0,
    name_template: str = "",
    surname_template: str = ""
) -> list[dict]:
    """
    Belirlenen aralıkta sıralı/seri barkod ve kod listesi oluşturur.
    Örn: prefix="URUN-", start_num=1, count=50, pad_length=4 -> URUN-0001 ... URUN-0050
    """
    items = []
    for i in range(count):
        cur_num = start_num + i
        num_str = str(cur_num)
        if pad_length > len(num_str):
            num_str = num_str.zfill(pad_length)
            
        full_code = f"{prefix}{num_str}{suffix}"
        
        name = name_template
        if "{no}" in name:
            name = name.replace("{no}", str(cur_num))
            
        surname = surname_template
        if "{no}" in surname:
            surname = surname.replace("{no}", str(cur_num))
            
        items.append({
            "name": name.strip(),
            "surname": surname.strip(),
            "tckn": full_code
        })
    return items
